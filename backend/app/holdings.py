"""Holdings: manual add/edit, and import from a broker export (Excel .xlsx or CSV)."""
from __future__ import annotations

import csv
import io
from decimal import Decimal, InvalidOperation

from sqlalchemy import select
from sqlalchemy.orm import Session

from .currencies import guess_currency
from .models import CURRENCIES, Broker, Category, Holding, Instrument
from .services import audit, row_dict



class ValidationError(ValueError):
    def __init__(self, message: str, field: str | None = None):
        super().__init__(message)
        self.field = field


def to_decimal(v, field: str, positive=False, allow_zero=True) -> Decimal | None:
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return None
    try:
        s = str(v).replace(",", "").strip()
        for sym in ("RM", "US$", "S$", "HK$", "$", "£"):
            s = s.replace(sym, "")
        d = Decimal(s.strip())
    except InvalidOperation as exc:
        raise ValidationError(f"{field}: '{v}' is not a number", field) from exc
    if not d.is_finite() or d < 0 or (positive and d == 0) or (not allow_zero and d == 0):
        raise ValidationError(f"{field} must be a positive number", field)
    return d


def get_or_create_instrument(db: Session, symbol: str, currency: str | None = None, name: str | None = None,
                             category_id: int | None = None, kind: str | None = None) -> tuple[Instrument, bool]:
    sym = symbol.strip().upper()
    if not sym:
        raise ValidationError("symbol is required", "symbol")
    inst = db.scalar(select(Instrument).where(Instrument.symbol == sym))
    if inst:
        if currency and currency.upper() != inst.currency:
            raise ValidationError(f"{sym} is priced in {inst.currency}, not {currency.upper()}", "currency")
        return inst, False
    ccy = (currency or guess_currency(sym)).upper()
    if ccy not in CURRENCIES:
        raise ValidationError(f"unsupported currency {ccy}", "currency")
    if len(sym) > 32:
        raise ValidationError("symbol or fund name is too long (32 characters max)", "symbol")
    kind = kind or "stock"
    if kind not in ("stock", "etf", "reit", "fund", "other"):
        raise ValidationError("invalid type", "kind")
    # Funds (EPF, ASB, ASNB, unit trusts) have no market feed: you enter their price or NAV.
    inst = Instrument(symbol=sym, currency=ccy, name=name or None, kind=kind,
                      price_mode="manual" if kind == "fund" else "auto", category_id=category_id,
                      withholding_tax_rate=Decimal("0.3") if (ccy == "USD" and kind != "fund") else Decimal(0))
    db.add(inst)
    db.flush()
    audit(db, "instrument", inst.id, "create", after=inst)
    return inst, True


def save_holding(db: Session, data: dict, holding: Holding | None = None) -> Holding:
    """Create or update. Adding a stock a broker already holds updates that holding."""
    qty = to_decimal(data.get("quantity", holding.quantity if holding else None), "quantity", positive=True)
    if qty is None:
        raise ValidationError("quantity is required", "quantity")
    total = to_decimal(data.get("total_cost"), "total_cost")
    if total is not None:  # funds: "amount you put in" instead of a per-unit cost
        cost = total / qty
    else:
        cost = to_decimal(data.get("avg_cost", holding.avg_cost if holding else None), "avg_cost")
    if cost is None:
        raise ValidationError("average cost is required", "avg_cost")
    broker = db.get(Broker, data.get("broker_id") or (holding.broker_id if holding else None))
    if broker is None:
        raise ValidationError("choose a broker", "broker_id")
    if holding is None:
        cat = data.get("category_id")
        if cat and db.get(Category, cat) is None:
            raise ValidationError("unknown category", "category_id")
        if data.get("instrument_id"):
            inst = db.get(Instrument, data["instrument_id"])
            if inst is None:
                raise ValidationError("unknown instrument", "instrument_id")
        else:
            inst, _ = get_or_create_instrument(db, data.get("symbol") or "", data.get("currency"), data.get("name"), cat,
                                               data.get("kind"))
        if cat and inst.category_id is None:
            inst.category_id = cat
        holding = db.scalar(select(Holding).where(Holding.broker_id == broker.id, Holding.instrument_id == inst.id))
    else:
        inst = holding.instrument
    before = row_dict(holding) if holding else None
    if holding is None:
        holding = Holding(broker_id=broker.id, instrument_id=inst.id, quantity=qty, avg_cost=cost,
                          source=data.get("source", "manual"))
        db.add(holding)
    else:
        clash = db.scalar(select(Holding).where(Holding.broker_id == broker.id, Holding.instrument_id == inst.id,
                                                Holding.id != holding.id))
        if clash:
            raise ValidationError(f"{broker.name} already has a {inst.symbol} holding", "broker_id")
        holding.broker_id, holding.quantity, holding.avg_cost = broker.id, qty, cost
        holding.source = data.get("source", holding.source)
    db.flush()
    price = to_decimal(data.get("price"), "price", positive=True)
    if price is not None:  # e.g. a fund's current NAV, entered with the holding
        from datetime import date

        from .market.service import save_manual_price
        save_manual_price(db, inst, price, date.today())
    audit(db, "holding", holding.id, "update" if before else "create", before=before, after=holding)
    return holding


def delete_holding(db: Session, holding: Holding):
    audit(db, "holding", holding.id, "delete", before=holding)
    db.delete(holding)
    db.flush()


# ---------------------------------------------------------------- import

COLUMNS = ["broker", "symbol", "name", "quantity", "avg_cost", "currency", "category"]
ALIASES = {
    "broker": ["broker", "account", "brokerage", "platform"],
    "symbol": ["symbol", "ticker", "code", "stock code", "stock", "counter"],
    "name": ["name", "stock name", "company", "description"],
    "quantity": ["quantity", "qty", "shares", "units", "holding", "holdings", "position", "volume"],
    "avg_cost": ["avg_cost", "avg cost", "average cost", "avg price", "average price", "cost price", "cost/share",
                 "cost per share", "avg. cost", "purchase price"],
    "currency": ["currency", "ccy"],
    "category": ["category", "group"],
    "kind": ["kind", "type", "asset type", "asset class"],
}
KINDS = {"stock": "stock", "equity": "stock", "etf": "etf", "reit": "reit", "fund": "fund", "unit trust": "fund"}


def read_table(content: bytes, filename: str) -> tuple[list[str], list[list[str]]]:
    """Return header and rows as strings. Excel: first sheet; the header is the first row with 2+ cells."""
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise ValidationError(f"Could not read the Excel file: {exc}") from exc
        ws = wb.worksheets[0]
        raw = [["" if c is None else str(c).strip() for c in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
    elif filename.lower().endswith(".xls"):
        raise ValidationError("Old .xls files are not supported. In Excel choose File > Save As > .xlsx or CSV.")
    else:
        text = None
        for enc in ("utf-8-sig", "cp1252"):
            try:
                text = content.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ValidationError("File is not readable text")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        raw = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), dialect)]
    raw = [r for r in raw if any(r)]
    start = next((i for i, r in enumerate(raw) if sum(1 for c in r if c) >= 2), None)
    if start is None:
        raise ValidationError("No table found in the file")
    header = raw[start]
    width = len(header)
    rows = [(r + [""] * width)[:width] for r in raw[start + 1:]]
    return header, rows


def suggest_mapping(header: list[str]) -> dict[str, int | None]:
    lower = [h.lower().strip() for h in header]
    out = {}
    for field, names in ALIASES.items():
        out[field] = next((lower.index(n) for n in names if n in lower), None)
    return out


def process_import(db: Session, content: bytes, filename: str, broker_id: int | None, mapping: dict | None,
                   mode: str, commit: bool) -> dict:
    """mode 'replace': the file is the broker's full list (holdings not in it are removed).
    mode 'merge': add or update only the rows in the file. Preview callers roll back the session."""
    if mode not in ("replace", "merge"):
        raise ValidationError("mode must be replace or merge")
    header, rows = read_table(content, filename)
    mapping = {k: (int(v) if v not in (None, "") else None) for k, v in (mapping or suggest_mapping(header)).items()}
    for f in ("symbol", "quantity", "avg_cost"):
        if mapping.get(f) is None:
            raise ValidationError(f"Choose which column holds {f.replace('_', ' ')}")
    if broker_id is None and mapping.get("broker") is None:
        raise ValidationError("Choose a broker, or a column that names the broker")
    fixed_broker = db.get(Broker, broker_id) if broker_id else None
    if broker_id and fixed_broker is None:
        raise ValidationError("unknown broker")

    def cell(r, f):
        i = mapping.get(f)
        return r[i].strip() if i is not None and i < len(r) else ""

    brokers = {b.name.lower(): b for b in db.scalars(select(Broker))}
    categories = {c.name.lower(): c for c in db.scalars(select(Category))}
    results, seen, created = [], {}, {"instruments": [], "brokers": [], "categories": []}
    touched_brokers: set[int] = set()
    for n, r in enumerate(rows, start=2):
        res = {"line": n, "symbol": cell(r, "symbol"), "status": "ok", "message": None}
        if not any(r):
            continue
        try:
            broker = fixed_broker
            if broker is None:
                name = cell(r, "broker")
                if not name:
                    raise ValidationError("broker is empty")
                broker = brokers.get(name.lower())
                if broker is None:
                    broker = Broker(name=name)
                    db.add(broker)
                    db.flush()
                    brokers[name.lower()] = broker
                    created["brokers"].append(name)
            cat_id = None
            if cell(r, "category"):
                cat = categories.get(cell(r, "category").lower())
                if cat is None:
                    cat = Category(name=cell(r, "category"))
                    db.add(cat)
                    db.flush()
                    categories[cat.name.lower()] = cat
                    created["categories"].append(cat.name)
                cat_id = cat.id
            qty = to_decimal(cell(r, "quantity"), "quantity")
            if not qty:
                res.update(status="skipped", message="quantity is zero")
                results.append(res)
                continue
            inst, new = get_or_create_instrument(db, cell(r, "symbol"), cell(r, "currency") or None,
                                                 cell(r, "name") or None, cat_id,
                                                 KINDS.get((cell(r, "kind") or "").strip().lower()))
            if new:
                created["instruments"].append(inst.symbol)
            if cat_id and inst.category_id != cat_id:
                inst.category_id = cat_id
            key = (broker.id, inst.id)
            if key in seen:
                raise ValidationError(f"{inst.symbol} appears twice for {broker.name} (line {seen[key]})")
            seen[key] = n
            existing = db.scalar(select(Holding).where(Holding.broker_id == broker.id, Holding.instrument_id == inst.id))
            res["status"] = "update" if existing else "new"
            if existing and Decimal(existing.quantity) == qty and Decimal(existing.avg_cost) == to_decimal(cell(r, "avg_cost"), "avg_cost"):
                res["status"] = "unchanged"
            save_holding(db, {"broker_id": broker.id, "instrument_id": inst.id, "quantity": qty,
                              "avg_cost": cell(r, "avg_cost"), "source": "import"}, existing)
            res["broker"] = broker.name
            res["quantity"] = str(qty)
            touched_brokers.add(broker.id)
        except ValidationError as exc:
            res.update(status="error", message=str(exc))
        results.append(res)

    removed = []
    if mode == "replace":
        scope = {fixed_broker.id} if fixed_broker else touched_brokers
        keep = set(seen)
        for h in db.scalars(select(Holding).where(Holding.broker_id.in_(scope))):
            if (h.broker_id, h.instrument_id) not in keep:
                removed.append({"broker": h.broker.name, "symbol": h.instrument.symbol, "quantity": str(h.quantity)})
                delete_holding(db, h)

    errors = sum(1 for r in results if r["status"] == "error")
    summary = {
        "rows": len(results),
        "new": sum(1 for r in results if r["status"] == "new"),
        "updated": sum(1 for r in results if r["status"] == "update"),
        "unchanged": sum(1 for r in results if r["status"] == "unchanged"),
        "skipped": sum(1 for r in results if r["status"] == "skipped"),
        "errors": errors,
        "removed": removed,
        "created": created,
    }
    if commit and errors:
        raise ValidationError("Fix the rows with errors (or remove them from the file) before importing")
    return {"header": header, "mapping": mapping, "rows": results[:500], "summary": summary, "can_import": errors == 0}


def export_csv(db: Session) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(COLUMNS)
    q = select(Holding).join(Broker).join(Instrument).order_by(Broker.name, Instrument.symbol)
    for h in db.scalars(q):
        i = h.instrument
        w.writerow([h.broker.name, i.symbol, i.name or "", str(h.quantity), str(h.avg_cost), i.currency,
                    i.category.name if i.category else ""])
    return buf.getvalue()
